"""
Додава колона `dialect` (yes/no) во dataset/metadata.csv.
Дијалектот се означува по ВИДЕО (сите клипови од едно видео добиваат иста вредност).

Чекор 1 - направи листа на видеа со автоматски предлог:
    python label_dialect.py make
  -> создава dialect_videos.xlsx. Провери ја колоната `dialect`
     и поправи ја каде предлогот е погрешен (yes / no).

Чекор 2 - запиши ја колоната во датасетот:
    python label_dialect.py apply

Ако повторно ги пуштиш merge скриптите, пушти го и `apply` пак.
"""
import argparse
import csv
import re
import sys
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parent
STEM = "dialect_videos"

# зборови типични за дијалект / разговорен говор (само за ПРЕДЛОГ)
MARKERS = {"шо", "саа", "тва", "тее", "абе", "бе", "чоек", "бре", "куде", "нејќи",
           "нејќам", "ич", "тоо", "шта", "мори", "кере", "ќере", "глеј", "питам",
           "питувам", "ного", "арно", "поарно", "аресе", "мее", "нем", "уд", "га", "ма"}
WORD = re.compile(r"[а-шѓќљњџѕјa-z]+", re.IGNORECASE)

def read_table(path, required=()):
    """Чита .xlsx или .csv (без разлика дали Excel го зачувал со , ; или tab,
    во UTF-8 / UTF-16 / cp1251, па дури и ако сè било во една колона)."""
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).active
        rows = [["" if c is None else str(c) for c in r] for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(x.strip() for x in r)]
        header = [h.strip().lower() for h in rows[0]] if rows else []
        if not all(c in header for c in required):
            sys.exit(f"Не можам да ги најдам колоните {', '.join(required)} во {path.name}")
        return [dict(zip(header, r)) for r in rows[1:]]

    data = path.read_bytes()
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        text = data.decode("utf-16")
    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1251")
    lines = [l for l in text.splitlines() if l.strip()]
    while lines and not lines[0].lower().lstrip('"').startswith(("file_name", "video", "sep=")):
        lines.pop(0)  # ѓубре пред header-от
    if lines and lines[0].lower().startswith("sep="):
        lines.pop(0)

    for delim in ("\t", ";", ","):
        rows = list(csv.reader(lines, delimiter=delim))
        if not rows:
            continue
        if len(rows[0]) == 1 and any(d in rows[0][0] for d in ";,\t"):
            inner = [r[0] if r else "" for r in rows]  # сè во една колона -> одмотај
            for d2 in ("\t", ";", ","):
                rr = list(csv.reader(inner, delimiter=d2))
                if all(c in [h.strip().lower() for h in rr[0]] for c in required):
                    rows = rr
                    break
        header = [h.strip().lower() for h in rows[0]]
        if all(c in header for c in required):
            return [dict(zip(header, r)) for r in rows[1:]]
    sys.exit(f"Не можам да ги најдам колоните {', '.join(required)} во {path.name}.\n"
             f"Прв ред: {lines[0][:150] if lines else '(празен)'}")


def write_xlsx(path, header, rows, widths=None, link_col=None):
    """Excel фајл - нема проблеми со кирилица, , или ; на ниеден компјутер."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        sys.exit("Треба openpyxl: pip install openpyxl")
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    if link_col is not None:
        for row in ws.iter_rows(min_row=2):
            cell = row[link_col]
            if cell.value:
                cell.hyperlink = str(cell.value)
                cell.value = "▶ play"
                cell.font = Font(color="0563C1", underline="single")
    for i, w in enumerate(widths or []):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"
    wb.save(path)


def find_list(stem):
    """Ја бара листата: прво .xlsx, па стар .csv."""
    for ext in (".xlsx", ".csv"):
        p = ROOT / (stem + ext)
        if p.exists():
            return p
    return None


def make(meta):
    existing = find_list(STEM)
    if existing:
        sys.exit(f"{existing.name} веќе постои - нема да го пребришам (избриши го рачно ако сакаш нов).")
    texts = defaultdict(list)
    for r in read_table(meta, ("video", "corrected_text")):
        texts[r["video"]].append(r["corrected_text"])
    out = []
    for video, ts in texts.items():
        full = " ".join(ts)
        words = WORD.findall(full.lower())
        score = sum(w in MARKERS for w in words) / max(len(words), 1)
        guess = "yes" if score >= 0.02 or "баба руж" in full.lower() else "no"
        out.append([video, guess, len(ts), round(score, 3), full[:250]])
    out.sort(key=lambda r: -r[3])
    path = ROOT / (STEM + ".xlsx")
    write_xlsx(path, ["video", "dialect", "clips", "score", "sample_text"], out,
               widths=[26, 9, 7, 7, 120])
    n = sum(r[1] == "yes" for r in out)
    print(f"Создаден {path.name}: {len(out)} видеа, предлог {n} yes / {len(out) - n} no.")
    print(f"Отвори го (start {path.name}), провери ја колоната `dialect`, зачувај, "
          f"па пушти: python label_dialect.py apply")


def apply(meta):
    path = find_list(STEM)
    if not path:
        sys.exit(f"Нема {STEM}.xlsx - прво пушти: python label_dialect.py make")
    labels = {}
    for r in read_table(path, ("video", "dialect")):
        v = (r.get("dialect") or "").strip().lower()
        v = {"da": "yes", "да": "yes", "не": "no", "ne": "no", "1": "yes", "0": "no"}.get(v, v)
        if v not in ("yes", "no"):
            sys.exit(f"Видео {r.get('video')}: dialect мора да е yes или no (сега е '{r.get('dialect')}')")
        labels[r["video"].strip()] = v

    rows = read_table(meta, ("file_name", "video"))
    missing = sorted({r["video"] for r in rows} - labels.keys())
    if missing:
        sys.exit(f"Овие видеа ги нема во {path.name}: {', '.join(missing)}")

    fields = [c for c in rows[0].keys() if c != "dialect"] + ["dialect"]
    for r in rows:
        r["dialect"] = labels[r["video"]]
    with open(meta, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)

    yes = [r for r in rows if r["dialect"] == "yes"]
    mins = lambda rs: sum(float(r["duration"] or 0) for r in rs) / 60
    print(f"Колона `dialect` додадена во {meta} (од {path.name})")
    print(f"  дијалект:    {len(yes)} клипови од {len({r['video'] for r in yes})} видеа ({mins(yes):.1f} мин)")
    print(f"  стандарден:  {len(rows) - len(yes)} клипови ({mins(rows) - mins(yes):.1f} мин)")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("step", choices=["make", "apply"])
    ap.add_argument("--dataset", default="dataset")
    a = ap.parse_args()
    meta = ROOT / a.dataset / "metadata.csv"
    if not meta.exists():
        sys.exit(f"Нема {meta}")
    make(meta) if a.step == "make" else apply(meta)
